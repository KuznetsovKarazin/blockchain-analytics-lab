#!/usr/bin/env python3
"""Build the two editable submission companions from reviewed LaTeX/Highlights.

Run with the primary artifact runtime Python (python-docx). These files do not
assert the absence of interests beyond the author's supplied funding disclosure.
"""
from pathlib import Path
import datetime
import io
import re
import zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

ARTICLE = Path(__file__).resolve().parents[1] / 'jss_latex_v8_final'


def plain(value):
    return value.replace('~', ' ').replace('\\%', '%').replace('\\&', '&')


def make_document(heading, title):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.left_margin = sec.right_margin = Inches(1)
    for style_name in ['Normal', 'Title', 'Subtitle', 'List Bullet']:
        style = doc.styles[style_name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        for fonts in style.element.xpath('.//w:rFonts'):
            for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
                fonts.attrib.pop(qn('w:' + attr), None)
    # The packaged Word template contains a decorative Title paragraph border.
    # Journal companion files use plain black typography with no title rule.
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    doc.styles['Title'].font.size = Pt(16)
    doc.styles['Title'].font.bold = True
    doc.styles['Normal'].paragraph_format.space_after = Pt(10)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.15
    doc.add_paragraph(heading, style='Title')
    para = doc.add_paragraph(title)
    para.paragraph_format.space_after = Pt(14)
    doc.add_paragraph('Oleksandr Kuznetsov')
    props = doc.core_properties
    props.author = 'Oleksandr Kuznetsov'
    props.title = heading
    props.subject = title
    props.last_modified_by = 'Oleksandr Kuznetsov'
    props.created = props.modified = datetime.datetime(2026, 10, 30, tzinfo=datetime.timezone.utc)
    props.revision = 1
    return doc


def stable_save(doc, path):
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    with zipfile.ZipFile(buf) as src, zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED, compresslevel=9) as dst:
        for name in sorted(src.namelist()):
            zi = zipfile.ZipInfo(name, date_time=(2026, 10, 30, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            dst.writestr(zi, src.read(name))


def main():
    tex = (ARTICLE / 'main.tex').read_text()
    title = plain(re.search(r'\\title\{([^}]+)\}', tex).group(1))
    bullets = [x[2:].strip() for x in (ARTICLE / 'HIGHLIGHTS.md').read_text().splitlines() if x.startswith('- ')]
    assert len(bullets) == 5 and all(len(x) <= 85 for x in bullets)
    doc = make_document('Highlights', title)
    for bullet in bullets:
        para = doc.add_paragraph(bullet, style='List Bullet')
        para.paragraph_format.space_after = Pt(12)
    stable_save(doc, ARTICLE / 'Highlights.docx')
    coi = plain(re.search(r'\\section\*\{Competing interests\}\s*(.*?)\s*\\section\*\{Funding\}', tex, re.S).group(1))
    assert '101308999' in coi and 'PETs4DeFi' in coi and '\\' not in coi
    doc = make_document('Declaration of competing interests', title)
    doc.add_paragraph(coi)
    stable_save(doc, ARTICLE / 'Declaration_of_Interest.docx')
    print('Created two submission companions; Highlights lengths:', [len(x) for x in bullets])


if __name__ == '__main__':
    main()
